"""Acme AI Q3 board update (FICTIONAL): writes the .pptx and an HTML twin (for the deck viewer's slide images).
# Run from backend/: uv run python samples/src/make_board_update.py samples/06_acme_board_update.pptx samples/src/acme_board_update.html
Text claims must stay: 22 full-time, 6 contractors in Lisbon, $12M raised, $25M Series A target opening November,
Globex (2 warehouses), Hooli, Initech LOI pilot Q1 2027; the KPI slide's speaker note says ARR closed Q3 at $2.6M."""
import sys, html
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT_PPTX, OUT_HTML = sys.argv[1], sys.argv[2]
NAVY, ORANGE, PAPER, MUTED, LINE, INK = "0F1B2D", "E0663A", "FBFAF7", "5D6878", "E6E2D9", "0F1B2D"
MONTHS = ["Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep"]
ARR = [0.99, 1.17, 1.35, 1.51, 1.72, 1.92, 2.13, 2.29, 2.42, 2.6]

SLIDES = [
    dict(kind="title", title="Acme AI - Q3 2026 board update", sub="October 2026 · Quarterly board meeting",
         note="Welcome. Agenda: highlights, metrics, team, financing, customers, then the asks."),
    dict(kind="bullets", k="Highlights", title="Q3 in one slide",
         bullets=["Globex went live in a second warehouse", "Plain-English explanations shipped to every site",
                  "Labour planning module in beta with Hooli", "Series A preparation underway"],
         note="The Globex expansion is the proof point for land-and-expand."),
    dict(kind="chart", k="Metrics", title="ARR by month",
         side=["Every customer expanded or renewed", "No churn this quarter", "Pipeline weighted to mid-market 3PLs"],
         note="ARR closed Q3 at $2.6M."),
    dict(kind="two", k="Team", title="Team", big="Headcount: 22 full-time (Sep 30)", body="Plus 6 contractors in Lisbon",
         right=["Hired: two engineers", "Open: Staff ML engineer", "Open: Head of Customer Success, Europe"],
         note="Hiring plan for Q4 is in the appendix of the pre-read."),
    dict(kind="two", k="Financing", title="Financing", big="Raised to date: $12M", body="Series A target: $25M, opening November",
         right=["October: data room and model ready", "November: first meetings", "December: target term sheet"],
         note="We'd like the board's help with warm intros to the firms on our list."),
    dict(kind="two", k="Customers", title="Customers", big="Globex (2 warehouses), Hooli", body="Initech: LOI signed, pilot Q1 2027",
         right=["Globex: expansion signed in August", "Hooli: labour module beta", "Initech: pilot scope agreed"],
         note="Initech's pilot is scoped to one building to start."),
    dict(kind="bullets", k="Asks", title="What we need from the board",
         bullets=["Intros to operations leaders at mid-market retailers", "Two references for the Head of Customer Success search",
                  "Feedback on the Series A target list before November"],
         note="Thank you. Questions?"),
]

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
prs.core_properties.title = "Acme AI Q3 2026 board update (FICTIONAL company, demo data)"
prs.core_properties.subject = "FICTIONAL company and data for the VentureSauce demo"
prs.core_properties.author = "Jane Founder"
W, H = prs.slide_width, prs.slide_height
px = lambda v: Emu(int(v * 12192000 / 1280))  # design in a 1280x720 grid


def rgb(h): return RGBColor.from_string(h)


def bg(slide, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid(); r.fill.fore_color.rgb = rgb(color); r.line.fill.background()


def text(slide, x, y, w, h, s, size, color=INK, bold=False, lines=None, spacing=1.15):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines or [s]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        r = p.add_run(); r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Helvetica Neue"
        r.font.color.rgb = rgb(color)
        if i: p.space_before = Pt(10)
    return tb


def footer(slide, n, dark=False):
    text(slide, 80, 676, 600, 24, "Acme AI · Board update · Confidential", 10, "9AA1AB")
    t = text(slide, 1100, 676, 100, 24, str(n), 10, "9AA1AB"); t.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


for n, d in enumerate(SLIDES, 1):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if d["kind"] == "title":
        bg(s, NAVY)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, px(80), px(70), px(22), px(22)); dot.fill.solid(); dot.fill.fore_color.rgb = rgb(ORANGE); dot.line.fill.background()
        text(s, 112, 64, 400, 36, "Acme AI", 18, "F4F1EA", True)
        text(s, 80, 260, 1000, 140, d["title"], 54, "F4F1EA", True)
        text(s, 80, 400, 900, 40, d["sub"], 20, "AEB6C4")
    else:
        bg(s, PAPER)
        text(s, 80, 56, 600, 24, d["k"].upper(), 12, ORANGE, True)
        text(s, 80, 84, 1000, 70, d["title"], 36, INK, True)
        if d["kind"] == "bullets":
            text(s, 80, 190, 1000, 400, "", 22, INK, lines=["•  " + b for b in d["bullets"]], spacing=1.3)
        elif d["kind"] == "chart":
            cd = CategoryChartData(); cd.categories = MONTHS; cd.add_series("ARR ($M)", ARR)
            ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, px(80), px(170), px(760), px(460), cd).chart
            ch.has_legend = False; ch.plots[0].gap_width = 60
            ser = ch.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(ORANGE)
            ch.value_axis.has_major_gridlines = False; ch.value_axis.visible = False
            ch.category_axis.tick_labels.font.size = Pt(11); ch.category_axis.tick_labels.font.color.rgb = rgb(MUTED)
            ch.category_axis.format.line.color.rgb = rgb(LINE)
            text(s, 880, 190, 320, 300, "", 18, INK, lines=["•  " + b for b in d["side"]], spacing=1.3)
        else:
            text(s, 80, 190, 620, 60, d["big"], 30, INK, True)
            text(s, 80, 250, 620, 40, d["body"], 22, MUTED)
            box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(760), px(180), px(440), px(300))
            box.fill.solid(); box.fill.fore_color.rgb = rgb("FFFFFF"); box.line.color.rgb = rgb(LINE); box.adjustments[0] = 0.06
            text(s, 790, 205, 390, 260, "", 17, INK, lines=["•  " + b for b in d["right"]], spacing=1.3)
        footer(s, n)
    s.notes_slide.notes_text_frame.text = d["note"]
prs.save(OUT_PPTX)

# HTML twin with the same content, for the viewer's slide images.
e = html.escape
css = """@page{size:1280px 720px;margin:0}*{box-sizing:border-box;margin:0}body{font-family:"Helvetica Neue",Helvetica,Arial,sans-serif;color:#0f1b2d}
.s{width:1280px;height:720px;position:relative;overflow:hidden;background:#fbfaf7;padding:56px 80px;page-break-after:always}.s.dark{background:#0f1b2d;color:#f4f1ea}
.k{font-size:16px;letter-spacing:.14em;font-weight:700;color:#e0663a}h2{font-size:48px;letter-spacing:-.02em;margin-top:12px}
ul{list-style:none;padding:0;margin-top:44px}li{font-size:29px;line-height:1.3;margin-bottom:20px;padding-left:34px;position:relative}li::before{content:"";position:absolute;left:0;top:13px;width:12px;height:12px;border-radius:50%;background:#e0663a}
.big{font-size:40px;font-weight:700;margin-top:48px;max-width:620px;line-height:1.15}.body{font-size:29px;color:#5d6878;margin-top:16px;max-width:620px}
.box{position:absolute;left:760px;top:180px;width:440px;height:300px;background:#fff;border:1px solid #e6e2d9;border-radius:18px;padding:26px 30px}.box li{font-size:22px;margin-bottom:16px;padding-left:26px}.box li::before{top:9px;width:9px;height:9px}.box ul{margin-top:0}
.foot{position:absolute;left:80px;right:80px;bottom:28px;display:flex;justify-content:space-between;font-size:13px;color:#9aa1ab}
.chart{position:absolute;left:80px;top:190px;width:760px;height:400px;display:flex;align-items:flex-end;gap:22px;border-bottom:2px solid #e6e2d9}.chart i{flex:1;background:#e0663a;border-radius:6px 6px 0 0}
.ax{position:absolute;left:80px;top:600px;width:760px;display:flex;gap:22px}.ax span{flex:1;text-align:center;font-size:14px;color:#5d6878}
.side{position:absolute;left:880px;top:200px;width:320px}.side li{font-size:22px}"""
out = [f"<!doctype html><html><head><meta charset=utf-8><title>Acme AI Q3 board update (FICTIONAL)</title><style>{css}</style></head><body>"]
for n, d in enumerate(SLIDES, 1):
    foot = f'<div class="foot"><span>Acme AI · Board update · Confidential</span><span>{n}</span></div>'
    if d["kind"] == "title":
        out.append(f'<section class="s dark"><div style="display:flex;gap:12px;align-items:center;font-weight:700;font-size:24px"><span style="width:26px;height:26px;border-radius:50%;background:#e0663a"></span>Acme AI</div>'
                   f'<h2 style="font-size:72px;margin-top:150px;max-width:1050px;line-height:1.02">{e(d["title"])}</h2><p style="font-size:26px;color:#aeb6c4;margin-top:24px">{e(d["sub"])}</p></section>')
        continue
    head = f'<section class="s"><p class="k">{e(d["k"].upper())}</p><h2>{e(d["title"])}</h2>'
    if d["kind"] == "bullets":
        body = "<ul>" + "".join(f"<li>{e(b)}</li>" for b in d["bullets"]) + "</ul>"
    elif d["kind"] == "chart":
        body = '<div class="chart">' + "".join(f'<i style="height:{v / max(ARR) * 100:.0f}%"></i>' for v in ARR) + '</div><div class="ax">' + "".join(f"<span>{m}</span>" for m in MONTHS) + "</div>"
        body += '<ul class="side">' + "".join(f"<li>{e(b)}</li>" for b in d["side"]) + "</ul>"
    else:
        body = f'<p class="big">{e(d["big"])}</p><p class="body">{e(d["body"])}</p><div class="box"><ul>' + "".join(f"<li>{e(b)}</li>" for b in d["right"]) + "</ul></div>"
    out.append(head + body + foot + "</section>")
out.append("</body></html>")
open(OUT_HTML, "w").write("\n".join(out))
print(len(SLIDES), "slides")
